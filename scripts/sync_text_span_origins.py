from __future__ import annotations

import argparse
import logging
import shutil
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


LOGGER = logging.getLogger(__name__)

SOURCE_REQUIRED_COLUMNS = ["id", "origin_x", "origin_y"]
BACKUP_FILENAME = "CAL-13311_translation_ready_before_origins.xlsx"


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Sync origin_x and origin_y values between text_spans sheets using id mapping."
    )
    parser.add_argument("source_xlsx", type=Path, help="Source workbook path (contains origin_x, origin_y)")
    parser.add_argument("target_xlsx", type=Path, help="Target workbook path to update")
    return parser.parse_args()


def to_abs(path_value: Path) -> Path:
    return path_value if path_value.is_absolute() else path_value.resolve()


def get_header_map(worksheet: Any) -> dict[str, int]:
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        raise RuntimeError(f"Worksheet '{worksheet.title}' has no header row")

    header_map: dict[str, int] = {}
    for index, value in enumerate(header_row):
        if value is None:
            continue
        header_map[str(value).strip()] = index
    return header_map


def col_to_letter(column_index_1_based: int) -> str:
    result = ""
    index = column_index_1_based
    while index > 0:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def parse_id(value: Any, row_index: int, workbook_label: str) -> int:
    if value is None:
        raise RuntimeError(f"Missing id at row {row_index} in {workbook_label}")
    try:
        return int(value)
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"Invalid id '{value}' at row {row_index} in {workbook_label}") from exc


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def load_source_map(source_ws: Any, source_label: str) -> tuple[dict[int, tuple[Any, Any]], list[int], set[int]]:
    header_map = get_header_map(source_ws)
    missing = [name for name in SOURCE_REQUIRED_COLUMNS if name not in header_map]
    if missing:
        raise RuntimeError(f"Missing required columns in source text_spans: {missing}")

    source_map: dict[int, tuple[Any, Any]] = {}
    duplicate_ids: set[int] = set()
    ids_without_origin: list[int] = []

    for row_idx, row in enumerate(source_ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        row_id = parse_id(row[header_map["id"]], row_idx, source_label)
        if row_id in source_map:
            duplicate_ids.add(row_id)
            continue

        origin_x = row[header_map["origin_x"]]
        origin_y = row[header_map["origin_y"]]
        if is_blank(origin_x) or is_blank(origin_y):
            ids_without_origin.append(row_id)

        source_map[row_id] = (origin_x, origin_y)

    return source_map, ids_without_origin, duplicate_ids


def ensure_target_origin_columns(target_ws: Any, target_header: dict[str, int]) -> tuple[int, int]:
    max_col = target_ws.max_column

    if "origin_x" not in target_header:
        max_col += 1
        target_ws.cell(row=1, column=max_col, value="origin_x")
        target_header["origin_x"] = max_col - 1

    if "origin_y" not in target_header:
        max_col += 1
        target_ws.cell(row=1, column=max_col, value="origin_y")
        target_header["origin_y"] = max_col - 1

    origin_x_col = target_header["origin_x"] + 1
    origin_y_col = target_header["origin_y"] + 1
    return origin_x_col, origin_y_col


def collect_target_rows(target_ws: Any, target_label: str, target_header: dict[str, int]) -> tuple[dict[int, int], set[int]]:
    if "id" not in target_header:
        raise RuntimeError("Missing required column 'id' in target text_spans")

    row_index_by_id: dict[int, int] = {}
    duplicate_ids: set[int] = set()

    for row_idx, row in enumerate(target_ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None for cell in row):
            continue

        row_id = parse_id(row[target_header["id"]], row_idx, target_label)
        if row_id in row_index_by_id:
            duplicate_ids.add(row_id)
            continue
        row_index_by_id[row_id] = row_idx

    return row_index_by_id, duplicate_ids


def sync_origins(source_xlsx: Path, target_xlsx: Path) -> int:
    source_abs = to_abs(source_xlsx)
    target_abs = to_abs(target_xlsx)

    if not source_abs.exists():
        raise RuntimeError(f"Source file not found: {source_abs}")
    if not target_abs.exists():
        raise RuntimeError(f"Target file not found: {target_abs}")

    source_wb = load_workbook(source_abs, data_only=True)
    target_wb = load_workbook(target_abs)

    if "text_spans" not in source_wb.sheetnames:
        raise RuntimeError("Worksheet 'text_spans' not found in source workbook")
    if "text_spans" not in target_wb.sheetnames:
        raise RuntimeError("Worksheet 'text_spans' not found in target workbook")

    source_ws = source_wb["text_spans"]
    target_ws = target_wb["text_spans"]

    source_map, ids_without_origin, source_duplicates = load_source_map(source_ws, "source workbook")
    target_header = get_header_map(target_ws)
    row_index_by_id, target_duplicates = collect_target_rows(target_ws, "target workbook", target_header)

    LOGGER.info("Source ids: %s", len(source_map))
    LOGGER.info("Target ids: %s", len(row_index_by_id))

    if source_duplicates:
        duplicate_list = sorted(source_duplicates)
        LOGGER.error("Duplicate ids in source: %s", duplicate_list)
        return 1

    if target_duplicates:
        duplicate_list = sorted(target_duplicates)
        LOGGER.error("Duplicate ids in target: %s", duplicate_list)
        return 1

    missing_in_target = sorted([row_id for row_id in source_map if row_id not in row_index_by_id])
    LOGGER.info("Missing ids in target: %s", missing_in_target)
    LOGGER.info("Ids without origin_x or origin_y: %s", sorted(ids_without_origin))

    if len(missing_in_target) > 5:
        LOGGER.error("More than five ids are missing in target (%s). Aborting without overwrite.", len(missing_in_target))
        return 1

    origin_x_col, origin_y_col = ensure_target_origin_columns(target_ws, target_header)

    updated_rows = 0
    for row_id, (origin_x, origin_y) in source_map.items():
        target_row_idx = row_index_by_id.get(row_id)
        if target_row_idx is None:
            continue

        target_ws.cell(row=target_row_idx, column=origin_x_col, value=origin_x)
        target_ws.cell(row=target_row_idx, column=origin_y_col, value=origin_y)
        updated_rows += 1

    backup_path = target_abs.parent / BACKUP_FILENAME
    shutil.copy2(target_abs, backup_path)
    target_wb.save(target_abs)

    LOGGER.info("Rows updated: %s", updated_rows)
    LOGGER.info("Backup created: %s", backup_path)
    LOGGER.info(
        "Target origin columns: %s, %s",
        col_to_letter(origin_x_col),
        col_to_letter(origin_y_col),
    )

    return 0


def main() -> int:
    setup_logging()
    args = parse_args()

    try:
        return sync_origins(args.source_xlsx, args.target_xlsx)
    except Exception as exc:  # noqa: BLE001
        LOGGER.error("Synchronization failed: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
