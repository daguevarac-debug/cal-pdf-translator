from __future__ import annotations

import argparse
import collections
import json
import logging
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


LOGGER = logging.getLogger(__name__)


@dataclass(slots=True)
class SpanRecord:
    id: int
    page: int
    original_text: str
    english_translation: str
    status: str
    bbox_x0: float
    bbox_y0: float
    bbox_x1: float
    bbox_y1: float
    font: str
    font_size: float
    color: int
    flags: int
    block_number: int
    line_number: int
    origin_x: float
    origin_y: float


@dataclass(slots=True)
class ExtractionSummary:
    pages: int
    blocks: int
    lines: int
    spans: int
    fonts: set[str]
    pages_without_text: list[int]


@dataclass(slots=True)
class LogicalLine:
    page: int
    block_number: int
    line_number: int
    segment_number: int
    text: str
    bbox_x0: float
    bbox_y0: float
    bbox_x1: float
    bbox_y1: float
    font: str
    font_size: float
    color: int
    flags: int
    source_span_ids: list[int]


@dataclass(slots=True)
class TranslationUnit:
    unit_id: int
    page: int
    action: str
    original_text: str
    english_translation: str
    bbox_x0: float
    bbox_y0: float
    bbox_x1: float
    bbox_y1: float
    font: str
    font_size: float
    color: int
    flags: int
    source_span_ids: list[int]
    block_number: int
    line_count: int


SPANISH_KEYWORDS = {
    "de",
    "del",
    "la",
    "el",
    "los",
    "las",
    "para",
    "por",
    "con",
    "sin",
    "una",
    "un",
    "se",
    "que",
    "en",
    "y",
    "certificado",
    "calibracion",
    "calibracion",
    "calibración",
    "medicion",
    "medición",
    "incertidumbre",
    "trazabilidad",
    "equipo",
    "numero",
    "número",
    "fecha",
    "rango",
    "valor",
    "promedio",
    "sesgo",
    "notas",
    "observaciones",
    "laboratorio",
    "cliente",
    "fabricante",
    "modelo",
    "temperatura",
    "humedad",
    "resultados",
    "autorizado",
    "recepcion",
    "recepción",
    "emision",
    "emisión",
    "pagina",
    "página",
    "codigo",
    "código",
    "documento",
    "restringido",
    "maximo",
    "máximo",
    "minima",
    "mínima",
    "relacion",
    "relación",
    "pedido",
    "direccion",
    "dirección",
    "ciudad",
    "serie",
    "tipo",
    "interno",
    "especificado",
    "medido",
    "expandida",
    "exactitud",
    "condiciones",
    "ambientales",
    "cobertura",
    "calibrador",
    "formato",
    "base",
}

ENGLISH_HINT_WORDS = {
    "certificate",
    "calibration",
    "measurement",
    "uncertainty",
    "traceability",
    "date",
    "model",
    "serial",
    "page",
    "temperature",
    "humidity",
    "results",
    "authorized",
    "document",
    "code",
    "address",
    "city",
    "laboratory",
    "client",
    "trademark",
    "licensed",
    "street",
    "city",
}

NO_SPACE_BEFORE = {",", ".", ":", ";", "%", ")"}

REMOVE_COLORS = {10921638, 9868950}

NUMERAL_START_RE = re.compile(r"^\s*\d+\.")
EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", flags=re.IGNORECASE)

EXPLICIT_TRANSLATE_PATTERNS = [
    re.compile(r"^\s*serie\s*$", flags=re.IGNORECASE),
    re.compile(r"^\s*tipo\s*$", flags=re.IGNORECASE),
    re.compile(r"^\s*interno\s*$", flags=re.IGNORECASE),
    re.compile(r"\bno\s+tiene\b", flags=re.IGNORECASE),
    re.compile(r"\bcondiciones\s+ambientales\b", flags=re.IGNORECASE),
    re.compile(r"\bcalibrador\s+ttr\b", flags=re.IGNORECASE),
    re.compile(r"\bformato\s+base\s*[:\-]?\s*[A-Z0-9][A-Z0-9\-_/\.]*\b", flags=re.IGNORECASE),
    re.compile(r"\bespecificado\b", flags=re.IGNORECASE),
    re.compile(r"\bmedido\b", flags=re.IGNORECASE),
    re.compile(r"\bexpandida\b", flags=re.IGNORECASE),
    re.compile(r"\bexactitud\b", flags=re.IGNORECASE),
    re.compile(r"\bcobertura\s*k\b", flags=re.IGNORECASE),
    re.compile(r"autopista\s+medell[ií]n\s+km\s*8,5\s*-\s*costado\s+sur", flags=re.IGNORECASE),
]

EXPLICIT_KEEP_VALUES = {
    "tenjo",
    "tenjo - cundinamarca",
    "copyright © siemens energy 2022, all rights reserved",
    "siemens energy is a trademark licensed by siemens ag",
    "tr-mark iii",
}

SIGNATURE_TOKENS = ("OID", "SERIALNUMBER", "CN=", "SN=", "G=", "C=", "O=", "OU=")

SECTION_HEADERS = {
    "datos del equipo",
    "información del cliente",
    "informacion del cliente",
    "método de calibración",
    "metodo de calibracion",
    "condiciones ambientales",
    "incertidumbre de medición",
    "incertidumbre de medicion",
    "trazabilidad",
    "resultados de calibración",
    "resultados de calibracion",
    "notas",
    "observaciones",
    "certificado de calibración",
    "certificado de calibracion",
}

ATOMIC_NON_MERGE_TEXTS = {
    "siemens energy is a trademark licensed by siemens ag",
    "laboratorio de calibración",
    "laboratorio de calibracion",
    "autopista medellín km 8,5 - costado sur",
    "autopista medellin km 8,5 - costado sur",
    "tenjo - cundinamarca",
    "tenjo",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inspect a PDF and export text spans, metadata, and page renders."
    )
    parser.add_argument("pdf_path", type=Path, help="Path to the source PDF file")
    return parser.parse_args()


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
    )


def ensure_directories(base_dir: Path) -> tuple[Path, Path]:
    data_dir = base_dir / "data"
    renders_dir = base_dir / "renders"
    data_dir.mkdir(parents=True, exist_ok=True)
    renders_dir.mkdir(parents=True, exist_ok=True)
    return data_dir, renders_dir


def approximate_origin(span: dict[str, Any], bbox: list[float] | tuple[float, ...]) -> tuple[float, float]:
    origin = span.get("origin")
    if isinstance(origin, (list, tuple)) and len(origin) == 2:
        return float(origin[0]), float(origin[1])

    if len(bbox) != 4:
        return 0.0, 0.0

    x0, _y0, _x1, y1 = bbox
    return float(x0), float(y1)


def extract_spans(doc: fitz.Document) -> tuple[list[SpanRecord], ExtractionSummary]:
    records: list[SpanRecord] = []
    fonts: set[str] = set()
    total_blocks = 0
    total_lines = 0
    pages_without_text: list[int] = []
    record_id = 1

    for page_index in range(doc.page_count):
        page = doc.load_page(page_index)
        page_dict = page.get_text("dict")
        page_blocks = page_dict.get("blocks", [])
        page_has_text = False

        for block_number, block in enumerate(page_blocks):
            if block.get("type") != 0:
                continue

            total_blocks += 1
            lines = block.get("lines", [])

            for line_number, line in enumerate(lines):
                total_lines += 1
                spans = line.get("spans", [])

                for span in spans:
                    text = str(span.get("text", ""))
                    if text.strip() == "":
                        continue

                    bbox_raw = span.get("bbox", [0.0, 0.0, 0.0, 0.0])
                    if not isinstance(bbox_raw, (list, tuple)) or len(bbox_raw) != 4:
                        bbox = (0.0, 0.0, 0.0, 0.0)
                    else:
                        bbox = tuple(float(v) for v in bbox_raw)

                    origin_x, origin_y = approximate_origin(span, bbox)

                    font = str(span.get("font", ""))
                    fonts.add(font)
                    record = SpanRecord(
                        id=record_id,
                        page=page_index + 1,
                        original_text=text,
                        english_translation="",
                        status="PENDING",
                        bbox_x0=bbox[0],
                        bbox_y0=bbox[1],
                        bbox_x1=bbox[2],
                        bbox_y1=bbox[3],
                        font=font,
                        font_size=float(span.get("size", 0.0)),
                        color=int(span.get("color", 0)),
                        flags=int(span.get("flags", 0)),
                        block_number=block_number,
                        line_number=line_number,
                        origin_x=origin_x,
                        origin_y=origin_y,
                    )
                    records.append(record)
                    record_id += 1
                    page_has_text = True

        if not page_has_text:
            pages_without_text.append(page_index + 1)

    summary = ExtractionSummary(
        pages=doc.page_count,
        blocks=total_blocks,
        lines=total_lines,
        spans=len(records),
        fonts=fonts,
        pages_without_text=pages_without_text,
    )
    return records, summary


def normalize_font_family(font_name: str) -> str:
    cleaned = font_name.lower().strip()
    if "+" in cleaned:
        cleaned = cleaned.split("+", maxsplit=1)[1]
    cleaned = cleaned.replace(" ", "").replace("_", "").replace("-", "")
    for suffix in ("psmt", "mt", "std", "regular"):
        if cleaned.endswith(suffix):
            cleaned = cleaned[: -len(suffix)]
    return cleaned


def equivalent_font(font_a: str, font_b: str) -> bool:
    if font_a == font_b:
        return True
    return normalize_font_family(font_a) == normalize_font_family(font_b)


def merge_span_texts(spans: list[SpanRecord]) -> str:
    ordered = sorted(spans, key=lambda s: (s.bbox_x0, s.bbox_y0, s.id))
    merged = ""
    prev_span: SpanRecord | None = None

    for span in ordered:
        token = span.original_text.strip()
        if token == "":
            continue

        if merged == "":
            merged = token
            prev_span = span
            continue

        add_space = True
        if token[0] in NO_SPACE_BEFORE:
            add_space = False
        elif merged.endswith("(") or merged.endswith("/") or merged.endswith("-"):
            add_space = False
        elif prev_span is not None:
            gap = span.bbox_x0 - prev_span.bbox_x1
            gap_threshold = max(0.8, min(span.font_size, prev_span.font_size) * 0.18)
            if gap <= gap_threshold:
                add_space = False

        if add_space:
            merged = f"{merged} {token}"
        else:
            merged = f"{merged}{token}"
        prev_span = span

    return merged.strip()


def is_bold(font_name: str, flags: int) -> bool:
    return "bold" in font_name.lower() or bool(flags & (1 << 4))


def is_italic(font_name: str, flags: int) -> bool:
    lower = font_name.lower()
    return "italic" in lower or "oblique" in lower or bool(flags & (1 << 1))


def span_style_key(span: SpanRecord) -> tuple[int, str, bool, bool, float]:
    return (
        span.color,
        normalize_font_family(span.font),
        is_bold(span.font, span.flags),
        is_italic(span.font, span.flags),
        span.font_size,
    )


def build_logical_lines(records: list[SpanRecord]) -> list[LogicalLine]:
    grouped: dict[tuple[int, int, int], list[SpanRecord]] = collections.defaultdict(list)
    for record in records:
        grouped[(record.page, record.block_number, record.line_number)].append(record)

    lines: list[LogicalLine] = []
    for (page, block_number, line_number), spans in grouped.items():
        ordered = sorted(spans, key=lambda s: (s.bbox_x0, s.id))

        style_segments: list[list[SpanRecord]] = []
        current_segment: list[SpanRecord] = []

        for span in ordered:
            if not current_segment:
                current_segment = [span]
                continue

            prev = current_segment[-1]
            prev_key = span_style_key(prev)
            curr_key = span_style_key(span)
            style_changed = (
                prev_key[0] != curr_key[0]
                or prev_key[1] != curr_key[1]
                or prev_key[2] != curr_key[2]
                or prev_key[3] != curr_key[3]
                or abs(prev_key[4] - curr_key[4]) > 0.5
            )

            if style_changed:
                style_segments.append(current_segment)
                current_segment = [span]
            else:
                current_segment.append(span)

        if current_segment:
            style_segments.append(current_segment)

        for segment_number, segment_spans in enumerate(style_segments):
            text = merge_span_texts(segment_spans)
            if text == "":
                continue

            bbox_x0 = min(s.bbox_x0 for s in segment_spans)
            bbox_y0 = min(s.bbox_y0 for s in segment_spans)
            bbox_x1 = max(s.bbox_x1 for s in segment_spans)
            bbox_y1 = max(s.bbox_y1 for s in segment_spans)

            weighted_font = collections.Counter[str]()
            weighted_color = collections.Counter[int]()
            weighted_flags = collections.Counter[int]()
            total_chars = 0
            weighted_font_size = 0.0

            for s in segment_spans:
                chars = max(1, len(s.original_text.strip()))
                total_chars += chars
                weighted_font[s.font] += chars
                weighted_color[s.color] += chars
                weighted_flags[s.flags] += chars
                weighted_font_size += s.font_size * chars

            dominant_font = weighted_font.most_common(1)[0][0]
            dominant_color = weighted_color.most_common(1)[0][0]
            dominant_flags = weighted_flags.most_common(1)[0][0]
            avg_font_size = weighted_font_size / max(1, total_chars)

            lines.append(
                LogicalLine(
                    page=page,
                    block_number=block_number,
                    line_number=line_number,
                    segment_number=segment_number,
                    text=text,
                    bbox_x0=bbox_x0,
                    bbox_y0=bbox_y0,
                    bbox_x1=bbox_x1,
                    bbox_y1=bbox_y1,
                    font=dominant_font,
                    font_size=avg_font_size,
                    color=dominant_color,
                    flags=dominant_flags,
                    source_span_ids=[s.id for s in segment_spans],
                )
            )

    lines.sort(
        key=lambda line: (
            line.page,
            line.block_number,
            line.line_number,
            line.segment_number,
            line.bbox_y0,
            line.bbox_x0,
        )
    )
    return lines


def line_width(line: LogicalLine) -> float:
    return line.bbox_x1 - line.bbox_x0


def in_same_column(left: LogicalLine, right: LogicalLine) -> bool:
    left_center = (left.bbox_x0 + left.bbox_x1) / 2.0
    right_center = (right.bbox_x0 + right.bbox_x1) / 2.0
    left_width = max(1.0, line_width(left))
    right_width = max(1.0, line_width(right))
    max_center_delta = min(80.0, 0.35 * max(left_width, right_width))
    return abs(left_center - right_center) <= max_center_delta


def should_merge_short_labels(prev_line: LogicalLine, next_line: LogicalLine) -> bool:
    if prev_line.page != next_line.page or prev_line.block_number != next_line.block_number:
        return False
    if next_line.line_number != prev_line.line_number + 1:
        return False
    if prev_line.color != next_line.color:
        return False
    if abs(prev_line.bbox_x0 - next_line.bbox_x0) >= 10.0:
        return False
    if not in_same_column(prev_line, next_line):
        return False
    vertical_gap = max(0.0, next_line.bbox_y0 - prev_line.bbox_y1)
    if vertical_gap >= (1.6 * max(prev_line.font_size, next_line.font_size)):
        return False
    if line_width(prev_line) >= 220.0 or line_width(next_line) >= 220.0:
        return False
    return True


def is_section_header(text: str) -> bool:
    normalized = normalize_text(text)
    return normalized in SECTION_HEADERS


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip().lower())


def is_atomic_non_merge_text(text: str) -> bool:
    normalized = normalize_text(text)
    if normalized in ATOMIC_NON_MERGE_TEXTS:
        return True
    if normalized.startswith("formato base:"):
        return True
    return False


def line_is_remove(line: LogicalLine) -> bool:
    return line.color in REMOVE_COLORS


def line_is_clearly_english(line: LogicalLine) -> bool:
    return looks_english(line.text) and not looks_spanish(line.text)


def line_style_compatible(prev_line: LogicalLine, next_line: LogicalLine) -> bool:
    if prev_line.color != next_line.color:
        return False
    if is_bold(prev_line.font, prev_line.flags) != is_bold(next_line.font, next_line.flags):
        return False
    if is_italic(prev_line.font, prev_line.flags) != is_italic(next_line.font, next_line.flags):
        return False
    if abs(prev_line.font_size - next_line.font_size) > 0.5:
        return False
    return True


def line_pair_blockers(prev_line: LogicalLine, next_line: LogicalLine) -> bool:
    if line_is_remove(prev_line) != line_is_remove(next_line):
        return True
    if looks_spanish(prev_line.text) and line_is_clearly_english(next_line):
        return True
    if looks_spanish(next_line.text) and line_is_clearly_english(prev_line):
        return True
    if is_section_header(prev_line.text):
        return True
    if is_section_header(next_line.text):
        return True
    if is_atomic_non_merge_text(prev_line.text) or is_atomic_non_merge_text(next_line.text):
        return True
    return False


def should_merge_lines(prev_line: LogicalLine, next_line: LogicalLine) -> bool:
    if prev_line.page != next_line.page or prev_line.block_number != next_line.block_number:
        return False
    if next_line.line_number != prev_line.line_number + 1:
        return False
    if prev_line.color != next_line.color:
        return False
    if not equivalent_font(prev_line.font, next_line.font):
        return False
    if abs(prev_line.font_size - next_line.font_size) >= 0.5:
        return False
    if abs(prev_line.bbox_x0 - next_line.bbox_x0) >= 6.0:
        return False

    vertical_gap = max(0.0, next_line.bbox_y0 - prev_line.bbox_y1)
    if vertical_gap >= (1.8 * max(prev_line.font_size, next_line.font_size)):
        return False
    if line_width(prev_line) <= 220.0 and line_width(next_line) <= 220.0:
        return False
    return True


def can_merge_lines(prev_line: LogicalLine, next_line: LogicalLine) -> bool:
    if NUMERAL_START_RE.match(next_line.text):
        return False
    if not line_style_compatible(prev_line, next_line):
        return False
    if line_pair_blockers(prev_line, next_line):
        return False
    return should_merge_lines(prev_line, next_line) or should_merge_short_labels(prev_line, next_line)


def merge_line_text(left: str, right: str) -> str:
    if left == "":
        return right
    if right == "":
        return left
    if right[0] in NO_SPACE_BEFORE:
        return f"{left}{right}"
    if left.endswith("(") or left.endswith("/") or left.endswith("-"):
        return f"{left}{right}"
    return f"{left} {right}"


def looks_spanish(text: str) -> bool:
    lowered = text.lower()
    if re.search(r"[áéíóúñü]", lowered):
        return True

    words = re.findall(r"[a-záéíóúñü]+", lowered)
    if not words:
        return False

    return any(word in SPANISH_KEYWORDS for word in words)


def spanish_words_in_text(text: str) -> set[str]:
    words = set(re.findall(r"[a-záéíóúñü]+", text.lower()))
    return words.intersection(SPANISH_KEYWORDS)


def has_code_like_value(text: str) -> bool:
    patterns = [
        r"\b[A-Z]{2,}-\d{2,}\b",
        r"\b\d{4,}\b",
        r"\b[A-Z0-9]{2,}[\-/][A-Z0-9\-/]{2,}\b",
        r"\b[A-Z]{1,3}\d{2,}[A-Z0-9-]*\b",
    ]
    return any(re.search(pattern, text) for pattern in patterns)


def is_spanish_label_with_code(text: str) -> bool:
    if ":" not in text:
        return False
    label, value = text.split(":", maxsplit=1)
    label = label.strip()
    value = value.strip()
    if label == "" or value == "":
        return False
    if not looks_spanish(label):
        return False
    return has_code_like_value(value)


def is_numeric_or_fixed_value(text: str) -> bool:
    cleaned = text.strip()
    if cleaned == "":
        return False

    regexes = [
        r"^[+-]?\d+(?:[\.,]\d+)?$",
        r"^[+-]?\d+(?:[\.,]\d+)?\s*%$",
        r"^\d{1,2}[\-/]\d{1,2}[\-/]\d{2,4}$",
        r"^\d{4}[\-/]\d{1,2}[\-/]\d{1,2}$",
        r"^[+-]?\d+(?:[\.,]\d+)?\s*(?:mm|cm|m|kg|g|mg|lb|hz|khz|mhz|ghz|v|mv|a|ma|ohm|kohm|mohm|°c|c|k|%rh|rpm|pa|kpa|mpa|bar|ms|s|min|h)$",
        r"^[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}$",
        r"^[A-Z0-9][A-Z0-9\-_/\.]{2,}$",
    ]

    for pattern in regexes:
        if re.match(pattern, cleaned, flags=re.IGNORECASE):
            return True

    return False


def looks_english(text: str) -> bool:
    words = re.findall(r"[a-zA-Z]+", text.lower())
    if not words:
        return False
    return any(word in ENGLISH_HINT_WORDS for word in words)


def is_explicit_translate(text: str) -> bool:
    return any(pattern.search(text) for pattern in EXPLICIT_TRANSLATE_PATTERNS)


def looks_like_proper_name(text: str) -> bool:
    cleaned = re.sub(r"\s+", " ", text.strip())
    if cleaned == "":
        return False
    if re.search(r"\d", cleaned):
        return False
    words = re.findall(r"[A-Za-z][A-Za-z\-']*", cleaned)
    if len(words) < 2:
        return False
    return all(word[0].isupper() for word in words)


def is_model_serial_or_code_without_spanish_label(text: str) -> bool:
    cleaned = text.strip()
    if cleaned == "":
        return False
    if is_spanish_label_with_code(cleaned):
        return False
    if re.match(r"^[A-Z0-9][A-Z0-9\-_/\.]{2,}$", cleaned):
        return True
    if re.match(r"^[A-Z]{1,5}[\- ]?[A-Z0-9]{2,}(?:[\-_/][A-Z0-9]+)*$", cleaned):
        return True
    return False


def is_terminal_expression(text: str) -> bool:
    cleaned = text.strip()
    if cleaned == "":
        return False
    if re.match(r"^H0-H[1-3]\s*/\s*X0-X[1-3]$", cleaned, flags=re.IGNORECASE):
        return True
    if re.match(r"^[A-Za-z0-9]+(?:[\-/][A-Za-z0-9]+)+$", cleaned):
        return True
    return False


def has_signature_token(text: str) -> bool:
    upper_text = text.upper()
    return any(token in upper_text for token in SIGNATURE_TOKENS)


def is_signature_line(line: LogicalLine) -> bool:
    if normalize_font_family(line.font) == "helvetica" and line.font_size <= 4.5:
        return True
    if has_signature_token(line.text):
        return True
    if EMAIL_RE.search(line.text):
        return True
    return False


def detect_signature_blocks(lines: list[LogicalLine]) -> set[tuple[int, int]]:
    blocks: set[tuple[int, int]] = set()
    for line in lines:
        if is_signature_line(line):
            blocks.add((line.page, line.block_number))
    return blocks


def classify_action(unit: TranslationUnit, signature_blocks: set[tuple[int, int]]) -> str:
    text = unit.original_text
    color = unit.color
    block_key = (unit.page, unit.block_number)
    lowered = text.strip().lower()

    if block_key in signature_blocks:
        return "KEEP"
    if is_explicit_translate(text):
        return "TRANSLATE"
    if lowered == "no":
        return "KEEP"
    if lowered in EXPLICIT_KEEP_VALUES:
        return "KEEP"
    if lowered.startswith("atrt check"):
        return "KEEP"
    if lowered.startswith("vanguard"):
        return "KEEP"
    if is_terminal_expression(text):
        return "KEEP"

    if color in REMOVE_COLORS:
        return "REMOVE"
    if is_spanish_label_with_code(text):
        return "TRANSLATE"
    if looks_spanish(text):
        return "TRANSLATE"
    if looks_like_proper_name(text):
        return "KEEP"
    if is_model_serial_or_code_without_spanish_label(text):
        return "KEEP"
    if is_numeric_or_fixed_value(text):
        return "KEEP"
    if looks_english(text):
        return "KEEP"
    return "REVIEW"


def build_translation_units(lines: list[LogicalLine]) -> list[TranslationUnit]:
    units: list[TranslationUnit] = []
    signature_blocks = detect_signature_blocks(lines)
    grouped: dict[tuple[int, int], list[LogicalLine]] = collections.defaultdict(list)
    for line in lines:
        grouped[(line.page, line.block_number)].append(line)

    unit_id = 1
    for (page, block_number), block_lines in sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1])):
        block_lines.sort(key=lambda line: (line.line_number, line.segment_number, line.bbox_y0, line.bbox_x0))
        current_lines: list[LogicalLine] = []

        for line in block_lines:
            if not current_lines:
                current_lines = [line]
                continue

            prev_line = current_lines[-1]
            if can_merge_lines(prev_line, line):
                current_lines.append(line)
            else:
                units.append(make_unit(current_lines, unit_id))
                unit_id += 1
                current_lines = [line]

        if current_lines:
            units.append(make_unit(current_lines, unit_id))
            unit_id += 1

    for unit in units:
        unit.action = classify_action(unit, signature_blocks)

    units = merge_short_headers_across_blocks(units, signature_blocks)
    units = merge_target_header_pairs(units, signature_blocks)
    for unit in units:
        unit.action = classify_action(unit, signature_blocks)

    return units


def unit_bold(unit: TranslationUnit) -> bool:
    return is_bold(unit.font, unit.flags)


def should_merge_across_blocks(first: TranslationUnit, second: TranslationUnit) -> bool:
    if first.page != second.page:
        return False
    if first.block_number == second.block_number:
        return False
    if first.action == "REMOVE" or second.action == "REMOVE":
        return False
    if is_atomic_non_merge_text(first.original_text) or is_atomic_non_merge_text(second.original_text):
        return False
    if normalize_text(first.original_text) in EXPLICIT_KEEP_VALUES:
        return False
    if normalize_text(second.original_text) in EXPLICIT_KEEP_VALUES:
        return False
    if (looks_spanish(first.original_text) and line_is_clearly_english_unit(second)) or (
        looks_spanish(second.original_text) and line_is_clearly_english_unit(first)
    ):
        return False
    if not (equivalent_font(first.font, "Arial-BoldMT") or unit_bold(first)):
        return False
    if not (equivalent_font(second.font, "Arial-BoldMT") or unit_bold(second)):
        return False
    if first.color != second.color:
        return False
    if abs(first.font_size - second.font_size) >= 0.5:
        return False
    if len(first.original_text) >= 40 or len(second.original_text) >= 40:
        return False
    if not re.search(r"[A-Za-zÁÉÍÓÚáéíóúÑñ]", first.original_text):
        return False
    if not re.search(r"[A-Za-zÁÉÍÓÚáéíóúÑñ]", second.original_text):
        return False
    if is_numeric_or_fixed_value(first.original_text) or is_numeric_or_fixed_value(second.original_text):
        return False
    if NUMERAL_START_RE.match(first.original_text) or NUMERAL_START_RE.match(second.original_text):
        return False

    first_h_center = (first.bbox_x0 + first.bbox_x1) / 2.0
    second_h_center = (second.bbox_x0 + second.bbox_x1) / 2.0
    overlap = max(0.0, min(first.bbox_x1, second.bbox_x1) - max(first.bbox_x0, second.bbox_x0))
    min_width = max(1.0, min(first.bbox_x1 - first.bbox_x0, second.bbox_x1 - second.bbox_x0))
    overlap_ratio = overlap / min_width
    if overlap_ratio <= 0.5 and abs(first_h_center - second_h_center) >= 12.0:
        return False

    vertical_gap = max(0.0, second.bbox_y0 - first.bbox_y1)
    if vertical_gap >= 1.5 * max(first.font_size, second.font_size):
        return False

    return True


def merge_two_units(first: TranslationUnit, second: TranslationUnit, new_id: int) -> TranslationUnit:
    text = merge_line_text(first.original_text, second.original_text)
    return TranslationUnit(
        unit_id=new_id,
        page=first.page,
        action="REVIEW",
        original_text=text,
        english_translation="",
        bbox_x0=min(first.bbox_x0, second.bbox_x0),
        bbox_y0=min(first.bbox_y0, second.bbox_y0),
        bbox_x1=max(first.bbox_x1, second.bbox_x1),
        bbox_y1=max(first.bbox_y1, second.bbox_y1),
        font=first.font,
        font_size=(first.font_size + second.font_size) / 2.0,
        color=first.color,
        flags=first.flags,
        source_span_ids=sorted(first.source_span_ids + second.source_span_ids),
        block_number=first.block_number,
        line_count=first.line_count + second.line_count,
    )


def merge_short_headers_across_blocks(
    units: list[TranslationUnit],
    signature_blocks: set[tuple[int, int]],
) -> list[TranslationUnit]:
    sorted_units = sorted(units, key=lambda u: (u.page, u.bbox_y0, u.bbox_x0, u.unit_id))
    merged: list[TranslationUnit] = []
    used_indices: set[int] = set()
    next_id = 1

    for index, current in enumerate(sorted_units):
        if index in used_indices:
            continue

        current_key = (current.page, current.block_number)
        if current_key in signature_blocks:
            current.unit_id = next_id
            merged.append(current)
            next_id += 1
            continue

        best_match_index: int | None = None
        best_gap: float | None = None
        for candidate_index in range(index + 1, len(sorted_units)):
            if candidate_index in used_indices:
                continue
            candidate = sorted_units[candidate_index]
            if candidate.page != current.page:
                break
            if candidate.bbox_y0 - current.bbox_y1 > 1.8 * max(current.font_size, candidate.font_size):
                break
            if (candidate.page, candidate.block_number) in signature_blocks:
                continue
            if not should_merge_across_blocks(current, candidate):
                continue

            gap = max(0.0, candidate.bbox_y0 - current.bbox_y1)
            if best_gap is None or gap < best_gap:
                best_gap = gap
                best_match_index = candidate_index

        if best_match_index is not None:
            nxt = sorted_units[best_match_index]
            combined = merge_two_units(current, nxt, next_id)
            merged.append(combined)
            used_indices.add(index)
            used_indices.add(best_match_index)
            next_id += 1
            continue

        current.unit_id = next_id
        merged.append(current)
        used_indices.add(index)
        next_id += 1

    return merged


def line_is_clearly_english_unit(unit: TranslationUnit) -> bool:
    return looks_english(unit.original_text) and not looks_spanish(unit.original_text)


def is_target_header_start(text: str) -> bool:
    normalized = normalize_text(text)
    return normalized in {"valor", "valor promedio", "incertidumbre"}


def is_target_header_end(text: str) -> bool:
    normalized = normalize_text(text)
    return normalized in {"especificado", "medido", "expandida"}


def can_merge_target_header(first: TranslationUnit, second: TranslationUnit) -> bool:
    if first.page != second.page:
        return False
    if not is_target_header_start(first.original_text):
        return False
    if not is_target_header_end(second.original_text):
        return False
    if first.action == "REMOVE" or second.action == "REMOVE":
        return False
    if not (unit_bold(first) and unit_bold(second)):
        return False
    if first.color != second.color:
        return False
    if abs(first.font_size - second.font_size) >= 0.5:
        return False
    if NUMERAL_START_RE.match(first.original_text) or NUMERAL_START_RE.match(second.original_text):
        return False

    overlap = max(0.0, min(first.bbox_x1, second.bbox_x1) - max(first.bbox_x0, second.bbox_x0))
    min_width = max(1.0, min(first.bbox_x1 - first.bbox_x0, second.bbox_x1 - second.bbox_x0))
    overlap_ratio = overlap / min_width
    first_center = (first.bbox_x0 + first.bbox_x1) / 2.0
    second_center = (second.bbox_x0 + second.bbox_x1) / 2.0
    if overlap_ratio <= 0.5 and abs(first_center - second_center) >= 12.0:
        return False

    vertical_gap = max(0.0, second.bbox_y0 - first.bbox_y1)
    if vertical_gap >= 1.5 * max(first.font_size, second.font_size):
        return False
    return True


def merge_target_header_pairs(
    units: list[TranslationUnit],
    signature_blocks: set[tuple[int, int]],
) -> list[TranslationUnit]:
    sorted_units = sorted(units, key=lambda u: (u.page, u.bbox_y0, u.bbox_x0, u.unit_id))
    used: set[int] = set()
    merged: list[TranslationUnit] = []
    next_id = 1

    for index, current in enumerate(sorted_units):
        if index in used:
            continue
        if (current.page, current.block_number) in signature_blocks:
            current.unit_id = next_id
            merged.append(current)
            next_id += 1
            used.add(index)
            continue

        best_index: int | None = None
        best_gap: float | None = None
        for candidate_index in range(index + 1, len(sorted_units)):
            if candidate_index in used:
                continue
            candidate = sorted_units[candidate_index]
            if candidate.page != current.page:
                break
            if (candidate.page, candidate.block_number) in signature_blocks:
                continue
            if not can_merge_target_header(current, candidate):
                continue

            gap = max(0.0, candidate.bbox_y0 - current.bbox_y1)
            if best_gap is None or gap < best_gap:
                best_gap = gap
                best_index = candidate_index

        if best_index is not None:
            combined = merge_two_units(current, sorted_units[best_index], next_id)
            merged.append(combined)
            used.add(index)
            used.add(best_index)
            next_id += 1
            continue

        current.unit_id = next_id
        merged.append(current)
        used.add(index)
        next_id += 1

    return merged


def make_unit(lines: list[LogicalLine], unit_id: int) -> TranslationUnit:
    text = ""
    source_span_ids: list[int] = []
    for line in lines:
        text = merge_line_text(text, line.text)
        source_span_ids.extend(line.source_span_ids)

    bbox_x0 = min(line.bbox_x0 for line in lines)
    bbox_y0 = min(line.bbox_y0 for line in lines)
    bbox_x1 = max(line.bbox_x1 for line in lines)
    bbox_y1 = max(line.bbox_y1 for line in lines)

    chars = [max(1, len(line.text)) for line in lines]
    total_chars = sum(chars)
    font_size = sum(line.font_size * weight for line, weight in zip(lines, chars)) / max(1, total_chars)

    font_counter = collections.Counter[str]()
    color_counter = collections.Counter[int]()
    flags_counter = collections.Counter[int]()
    for line, weight in zip(lines, chars):
        font_counter[line.font] += weight
        color_counter[line.color] += weight
        flags_counter[line.flags] += weight

    font = font_counter.most_common(1)[0][0]
    color = color_counter.most_common(1)[0][0]
    flags = flags_counter.most_common(1)[0][0]
    action = "REVIEW"

    return TranslationUnit(
        unit_id=unit_id,
        page=lines[0].page,
        action=action,
        original_text=text,
        english_translation="",
        bbox_x0=bbox_x0,
        bbox_y0=bbox_y0,
        bbox_x1=bbox_x1,
        bbox_y1=bbox_y1,
        font=font,
        font_size=font_size,
        color=color,
        flags=flags,
        source_span_ids=sorted(source_span_ids),
        block_number=lines[0].block_number,
        line_count=len(lines),
    )


def find_suspicious_merged_units(units: list[TranslationUnit]) -> list[TranslationUnit]:
    suspicious: list[TranslationUnit] = []
    for unit in units:
        text = unit.original_text.strip()
        text_len = len(text)
        has_spanish = looks_spanish(text)
        has_aux_english = looks_english(text) and ("/" in text or unit.color in REMOVE_COLORS)
        unit_height = unit.bbox_y1 - unit.bbox_y0
        is_tall_short = unit_height > 4.0 * max(0.1, unit.font_size) and text_len < 180
        many_lines_non_paragraph = unit.line_count > 3 and not (text_len > 220 and line_width_from_bbox(unit) > 220)

        if many_lines_non_paragraph or (has_spanish and has_aux_english) or is_tall_short:
            suspicious.append(unit)
    return suspicious


def line_width_from_bbox(unit: TranslationUnit) -> float:
    return unit.bbox_x1 - unit.bbox_x0


def save_json(records: list[SpanRecord], summary: ExtractionSummary, json_path: Path) -> None:
    payload = {
        "summary": {
            "pages": summary.pages,
            "blocks": summary.blocks,
            "lines": summary.lines,
            "spans": summary.spans,
            "fonts": sorted(summary.fonts),
            "pages_without_text": summary.pages_without_text,
        },
        "spans": [asdict(record) for record in records],
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def save_excel(records: list[SpanRecord], units: list[TranslationUnit], xlsx_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "text_spans"

    headers = [
        "id",
        "page",
        "original_text",
        "english_translation",
        "status",
        "bbox_x0",
        "bbox_y0",
        "bbox_x1",
        "bbox_y1",
        "font",
        "font_size",
        "color",
        "flags",
        "block_number",
        "line_number",
        "origin_x",
        "origin_y",
    ]
    ws.append(headers)

    for record in records:
        ws.append(
            [
                record.id,
                record.page,
                record.original_text,
                record.english_translation,
                record.status,
                record.bbox_x0,
                record.bbox_y0,
                record.bbox_x1,
                record.bbox_y1,
                record.font,
                record.font_size,
                record.color,
                record.flags,
                record.block_number,
                record.line_number,
                record.origin_x,
                record.origin_y,
            ]
        )

    ws_units = wb.create_sheet("translation_units")
    unit_headers = [
        "unit_id",
        "page",
        "action",
        "original_text",
        "english_translation",
        "bbox_x0",
        "bbox_y0",
        "bbox_x1",
        "bbox_y1",
        "font",
        "font_size",
        "color",
        "flags",
        "source_span_ids",
        "block_number",
    ]
    ws_units.append(unit_headers)

    for unit in units:
        ws_units.append(
            [
                unit.unit_id,
                unit.page,
                unit.action,
                unit.original_text,
                unit.english_translation,
                unit.bbox_x0,
                unit.bbox_y0,
                unit.bbox_x1,
                unit.bbox_y1,
                unit.font,
                unit.font_size,
                unit.color,
                unit.flags,
                ",".join(str(span_id) for span_id in unit.source_span_ids),
                unit.block_number,
            ]
        )

    format_translation_units_sheet(ws_units)

    wb.save(xlsx_path)


def format_translation_units_sheet(worksheet: Any) -> None:
    header_font = Font(bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    action_fills = {
        "TRANSLATE": PatternFill(fill_type="solid", fgColor="FFFCE5CD"),
        "KEEP": PatternFill(fill_type="solid", fgColor="FFD9EAD3"),
        "REMOVE": PatternFill(fill_type="solid", fgColor="FFE6E6E6"),
        "REVIEW": PatternFill(fill_type="solid", fgColor="FFF4CCCC"),
    }

    for cell in worksheet[1]:
        cell.font = header_font

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    worksheet.column_dimensions["A"].width = 10
    worksheet.column_dimensions["B"].width = 8
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 70
    worksheet.column_dimensions["E"].width = 70
    worksheet.column_dimensions["F"].width = 11
    worksheet.column_dimensions["G"].width = 11
    worksheet.column_dimensions["H"].width = 11
    worksheet.column_dimensions["I"].width = 11
    worksheet.column_dimensions["J"].width = 20
    worksheet.column_dimensions["K"].width = 11
    worksheet.column_dimensions["L"].width = 12
    worksheet.column_dimensions["M"].width = 8
    worksheet.column_dimensions["N"].width = 24
    worksheet.column_dimensions["O"].width = 12

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = wrap_alignment
        action = str(row[2].value or "").upper()
        fill = action_fills.get(action)
        if fill is not None:
            row[2].fill = fill


def render_pages(doc: fitz.Document, render_root: Path) -> None:
    dpi = 150
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)

    for page_index in range(doc.page_count):
        page = doc.load_page(page_index)
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        output_path = render_root / f"page_{page_index + 1:03d}.png"
        pix.save(output_path)


def print_summary(summary: ExtractionSummary, units: list[TranslationUnit]) -> None:
    fonts_sorted = sorted(font for font in summary.fonts if font)
    pages_without_text = ", ".join(str(p) for p in summary.pages_without_text) or "None"
    action_counter = collections.Counter(unit.action for unit in units)

    print("PDF inspection summary")
    print(f"- Number of pages: {summary.pages}")
    print(f"- Text blocks: {summary.blocks}")
    print(f"- Text lines: {summary.lines}")
    print(f"- Text spans: {summary.spans}")
    print(f"- Fonts found: {', '.join(fonts_sorted) if fonts_sorted else 'None'}")
    print(f"- Pages without selectable text: {pages_without_text}")
    print(f"- Translation units: {len(units)}")
    print(
        "- Units by action: "
        f"TRANSLATE={action_counter.get('TRANSLATE', 0)}, "
        f"KEEP={action_counter.get('KEEP', 0)}, "
        f"REMOVE={action_counter.get('REMOVE', 0)}, "
        f"REVIEW={action_counter.get('REVIEW', 0)}"
    )
    suspicious = find_suspicious_merged_units(units)
    print("- suspicious_merged_units:")
    if not suspicious:
        print("  None")
    else:
        for unit in suspicious:
            preview = re.sub(r"\s+", " ", unit.original_text).strip()
            print(
                "  "
                f"unit_id={unit.unit_id} page={unit.page} block={unit.block_number} "
                f"lines={unit.line_count} len={len(preview)} text={preview[:120]}"
            )


def warn_review_units_with_spanish(units: list[TranslationUnit]) -> None:
    for unit in units:
        if unit.action != "REVIEW":
            continue
        matched_words = spanish_words_in_text(unit.original_text)
        if matched_words:
            LOGGER.warning(
                "REVIEW unit contains Spanish keywords | unit_id=%s | page=%s | block=%s | words=%s | text=%s",
                unit.unit_id,
                unit.page,
                unit.block_number,
                ",".join(sorted(matched_words)),
                unit.original_text,
            )


def main() -> int:
    setup_logging()
    args = parse_args()

    script_path = Path(__file__).resolve()
    project_root = script_path.parent.parent
    pdf_path: Path = args.pdf_path
    if not pdf_path.is_absolute():
        pdf_path = (project_root / pdf_path).resolve()

    if not pdf_path.exists():
        LOGGER.error("PDF file does not exist: %s", pdf_path)
        return 1
    if pdf_path.suffix.lower() != ".pdf":
        LOGGER.error("Input file is not a PDF: %s", pdf_path)
        return 1

    data_dir, renders_dir = ensure_directories(project_root)
    pdf_name = pdf_path.stem
    json_path = data_dir / f"{pdf_name}_text.json"
    xlsx_path = data_dir / f"{pdf_name}_text.xlsx"
    render_output_dir = renders_dir / pdf_name
    render_output_dir.mkdir(parents=True, exist_ok=True)

    try:
        with fitz.open(pdf_path) as doc:
            records, summary = extract_spans(doc)
            logical_lines = build_logical_lines(records)
            units = build_translation_units(logical_lines)
            save_json(records, summary, json_path)
            save_excel(records, units, xlsx_path)
            render_pages(doc, render_output_dir)

    except ImportError as exc:
        LOGGER.error("Missing dependency: %s", exc)
        return 1
    except fitz.FileDataError as exc:
        LOGGER.error("Could not open or parse PDF file: %s", exc)
        return 1
    except OSError as exc:
        LOGGER.error("Filesystem error: %s", exc)
        return 1
    except Exception as exc:  # noqa: BLE001
        LOGGER.exception("Unexpected error while processing PDF: %s", exc)
        return 1

    LOGGER.info("JSON written to: %s", json_path)
    LOGGER.info("Excel written to: %s", xlsx_path)
    LOGGER.info("Rendered pages written to: %s", render_output_dir)
    warn_review_units_with_spanish(units)
    print_summary(summary, units)
    return 0


if __name__ == "__main__":
    sys.exit(main())