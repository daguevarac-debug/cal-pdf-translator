from __future__ import annotations

import argparse
import json
import logging
import re
import statistics
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from openpyxl import load_workbook


LOGGER = logging.getLogger(__name__)

REQUIRED_COLUMNS = [
    "unit_id",
    "page",
    "action",
    "original_text",
    "english_translation",
    "bbox_x0",
    "bbox_y0",
    "bbox_x1",
    "bbox_y1",
    "font",
    "font_size",
    "color",
    "flags",
    "source_span_ids",
    "block_number",
]

TRANSLATION_NOTE = "Unofficial English translation - the original Spanish certificate prevails."


@dataclass(slots=True)
class TranslationUnitRow:
    unit_id: int
    page: int
    action: str
    original_text: str
    english_translation: str
    bbox_x0: float
    bbox_y0: float
    bbox_x1: float
    bbox_y1: float
    font: str
    font_size: float
    color: int
    flags: int
    source_span_ids: str
    block_number: int

    @property
    def bbox(self) -> fitz.Rect:
        return fitz.Rect(self.bbox_x0, self.bbox_y0, self.bbox_x1, self.bbox_y1)


@dataclass(slots=True)
class Summary:
    translate_processed: int = 0
    remove_processed: int = 0
    keep_ignored: int = 0
    empty_translations: int = 0
    reduced_font_texts: int = 0
    not_fit_texts: int = 0
    row_failures: int = 0


@dataclass(slots=True)
class FontSpec:
    fontname: str
    fontfile: str | None


@dataclass(slots=True)
class FitResult:
    fits: bool
    proposed_size: float


@dataclass(slots=True)
class TextSpanRow:
    span_id: int
    page: int
    bbox_x0: float
    bbox_y0: float
    bbox_x1: float
    bbox_y1: float
    font: str
    font_size: float
    flags: int
    block_number: int
    line_number: int
    origin_x: float
    origin_y: float

    @property
    def bbox(self) -> fitz.Rect:
        return fitz.Rect(self.bbox_x0, self.bbox_y0, self.bbox_x1, self.bbox_y1)


@dataclass(slots=True)
class PlacementSpec:
    placement_mode: str
    font_size_original: float
    font_size_final: float
    align: int
    use_textbox: bool
    lineheight: float | None = None
    forced_text: str | None = None


TEXT_SPAN_REQUIRED_COLUMNS = [
    "id",
    "page",
    "bbox_x0",
    "bbox_y0",
    "bbox_x1",
    "bbox_y1",
    "font",
    "font_size",
    "flags",
    "block_number",
    "line_number",
]

SECTION_TITLES = {
    "equipment information",
    "customer information",
    "calibration method",
    "environmental conditions",
    "measurement uncertainty",
    "traceability",
    "calibration results",
    "notes",
    "observations",
}

EXPLICIT_CENTER_TEXTS = {
    "calibration certificate",
    "ratio measurement",
    "end of certificate",
    "** end of certificate **",
}

PAGE2_TABLE_HEADER_TEXTS = {
    "equipment": "Equipment",
    "type": "Type",
    "internal number": "Internal\nNumber",
    "calibrated by": "Calibrated By",
    "certificate no.": "Certificate No.",
    "calibration date": "Calibration\nDate",
}

PAGE3_TABLE_HEADER_TEXTS = {
    "range": "Range",
    "specified value": "Specified\nValue",
    "average measured value": "Average\nMeasured\nValue",
    "bias": "Bias",
    "expanded uncertainty": "Expanded\nUncertainty",
    "coverage factor k": "Coverage\nFactor k",
}

SHORT_LEFT_LABELS = {
    "equipment information",
    "description",
    "manufacturer",
    "serial number",
    "model",
    "customer identification",
    "customer information",
    "customer name",
    "address",
    "city",
    "order number",
    "calibration method",
    "environmental conditions",
    "date received",
    "calibration date",
    "date of issue",
    "authorized by",
}


def normalize_text(value: str) -> str:
    collapsed = re.sub(r"\s+", " ", value.strip().lower())
    return collapsed.replace("\u00a0", " ")


def display_text_for_row(row: TranslationUnitRow) -> str:
    preferred = row.english_translation.strip()
    if preferred != "":
        return preferred
    return row.original_text.strip()


def parse_source_span_ids(raw_value: str) -> list[int]:
    return [int(match) for match in re.findall(r"\d+", raw_value or "")]


def is_section_title_row(row: TranslationUnitRow) -> bool:
    text = normalize_text(display_text_for_row(row))
    return text in SECTION_TITLES


def is_page2_table_header_row(row: TranslationUnitRow) -> bool:
    if row.page != 2:
        return False
    return normalize_text(display_text_for_row(row)) in PAGE2_TABLE_HEADER_TEXTS


def is_page3_table_header_row(row: TranslationUnitRow) -> bool:
    if row.page != 3:
        return False
    return normalize_text(display_text_for_row(row)) in PAGE3_TABLE_HEADER_TEXTS


def is_notes_row(row: TranslationUnitRow) -> bool:
    if row.page != 3:
        return False
    return re.match(r"^\s*[1-7]\.", row.original_text or "") is not None


def is_short_label_row(row: TranslationUnitRow) -> bool:
    text = display_text_for_row(row)
    normalized = normalize_text(text)
    if normalized in SHORT_LEFT_LABELS:
        return True
    if len(text) >= 80:
        return False
    if "\n" in text:
        return False
    if is_section_title_row(row) or is_page2_table_header_row(row) or is_page3_table_header_row(row):
        return False
    return len(text) > 0


def is_page2_paragraph_row(row: TranslationUnitRow) -> bool:
    if row.page != 2:
        return False
    if is_section_title_row(row) or is_page2_table_header_row(row) or is_notes_row(row):
        return False
    text = display_text_for_row(row)
    if len(text) <= 80:
        return False
    return True


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Apply reviewed translations onto a source PDF.")
    parser.add_argument("pdf_path", type=Path, help="Path to original Spanish PDF")
    parser.add_argument("excel_path", type=Path, help="Path to reviewed Excel file")
    parser.add_argument("--sheet", default="translation_ready", help="Worksheet name with translation units")
    parser.add_argument("--output", type=Path, default=None, help="Output PDF path")
    parser.add_argument("--dry-run", action="store_true", help="Run preflight checks only and generate report JSON")
    return parser.parse_args()


def to_project_path(project_root: Path, path_value: Path) -> Path:
    if path_value.is_absolute():
        return path_value
    return (project_root / path_value).resolve()


def normalize_font_family(font_name: str) -> str:
    cleaned = font_name.lower().strip()
    if "+" in cleaned:
        cleaned = cleaned.split("+", maxsplit=1)[1]
    cleaned = cleaned.replace(" ", "").replace("_", "").replace("-", "")
    for suffix in ("psmt", "mt", "std", "regular"):
        if cleaned.endswith(suffix):
            cleaned = cleaned[: -len(suffix)]
    return cleaned


def is_bold(font_name: str, flags: int) -> bool:
    return "bold" in font_name.lower() or bool(flags & (1 << 4))


def is_italic(font_name: str, flags: int) -> bool:
    lower = font_name.lower()
    return "italic" in lower or "oblique" in lower or bool(flags & (1 << 1))


def choose_font_spec(unit: TranslationUnitRow, arial_paths: dict[str, Path]) -> FontSpec:
    source_font = unit.font
    bold = is_bold(source_font, unit.flags)
    italic = is_italic(source_font, unit.flags)
    normalized = normalize_font_family(source_font)

    if normalized == "helvetica":
        if bold and italic:
            return FontSpec("hebi", None)
        if bold:
            return FontSpec("hebo", None)
        if italic:
            return FontSpec("heit", None)
        return FontSpec("helv", None)

    if bold and italic:
        key = "arialbi"
        fallback = "hebi"
    elif bold:
        key = "arialbd"
        fallback = "hebo"
    elif italic:
        key = "ariali"
        fallback = "heit"
    else:
        key = "arial"
        fallback = "helv"

    arial_path = arial_paths.get(key)
    if arial_path is not None and arial_path.exists():
        return FontSpec(fontname=f"f_{key}", fontfile=str(arial_path))
    return FontSpec(fontname=fallback, fontfile=None)


def parse_int(value: Any, field_name: str) -> int:
    if value is None:
        raise ValueError(f"Missing value for {field_name}")
    if isinstance(value, bool):
        raise ValueError(f"Invalid boolean for {field_name}")
    return int(value)


def parse_float(value: Any, field_name: str) -> float:
    if value is None:
        raise ValueError(f"Missing value for {field_name}")
    return float(value)


def parse_row(row_values: dict[str, Any]) -> TranslationUnitRow:
    action = str(row_values["action"] or "").strip().upper()
    return TranslationUnitRow(
        unit_id=parse_int(row_values["unit_id"], "unit_id"),
        page=parse_int(row_values["page"], "page"),
        action=action,
        original_text=str(row_values["original_text"] or ""),
        english_translation=str(row_values["english_translation"] or ""),
        bbox_x0=parse_float(row_values["bbox_x0"], "bbox_x0"),
        bbox_y0=parse_float(row_values["bbox_y0"], "bbox_y0"),
        bbox_x1=parse_float(row_values["bbox_x1"], "bbox_x1"),
        bbox_y1=parse_float(row_values["bbox_y1"], "bbox_y1"),
        font=str(row_values["font"] or ""),
        font_size=parse_float(row_values["font_size"], "font_size"),
        color=parse_int(row_values["color"], "color"),
        flags=parse_int(row_values["flags"], "flags"),
        source_span_ids=str(row_values["source_span_ids"] or ""),
        block_number=parse_int(row_values["block_number"], "block_number"),
    )


def read_translation_units(excel_path: Path, sheet_name: str, workbook: Any | None = None) -> list[TranslationUnitRow]:
    if workbook is None:
        try:
            workbook = load_workbook(excel_path, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError(f"Could not open Excel file: {excel_path} | {exc}") from exc

    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Worksheet '{sheet_name}' not found in {excel_path}")

    worksheet = workbook[sheet_name]
    header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_cells is None:
        raise RuntimeError(f"Worksheet '{sheet_name}' has no header row")

    header_map: dict[str, int] = {}
    for index, value in enumerate(header_cells):
        if value is None:
            continue
        header_map[str(value).strip()] = index

    missing_columns = [column for column in REQUIRED_COLUMNS if column not in header_map]
    if missing_columns:
        raise RuntimeError(f"Missing required columns in sheet '{sheet_name}': {missing_columns}")

    rows: list[TranslationUnitRow] = []
    for excel_row in worksheet.iter_rows(min_row=2, values_only=True):
        if all(cell is None for cell in excel_row):
            continue

        row_values = {column: excel_row[header_map[column]] for column in REQUIRED_COLUMNS}
        rows.append(parse_row(row_values))

    return rows


def read_text_spans(workbook: Any) -> dict[int, TextSpanRow]:
    if "text_spans" not in workbook.sheetnames:
        raise RuntimeError("Worksheet 'text_spans' not found in workbook")

    worksheet = workbook["text_spans"]
    header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_cells is None:
        raise RuntimeError("Worksheet 'text_spans' has no header row")

    header_map: dict[str, int] = {}
    for index, value in enumerate(header_cells):
        if value is None:
            continue
        header_map[str(value).strip()] = index

    missing_columns = [column for column in TEXT_SPAN_REQUIRED_COLUMNS if column not in header_map]
    if missing_columns:
        raise RuntimeError(f"Missing required columns in sheet 'text_spans': {missing_columns}")

    span_map: dict[int, TextSpanRow] = {}
    for excel_row in worksheet.iter_rows(min_row=2, values_only=True):
        if all(cell is None for cell in excel_row):
            continue

        span_id = parse_int(excel_row[header_map["id"]], "id")
        span = TextSpanRow(
            span_id=span_id,
            page=parse_int(excel_row[header_map["page"]], "page"),
            bbox_x0=parse_float(excel_row[header_map["bbox_x0"]], "bbox_x0"),
            bbox_y0=parse_float(excel_row[header_map["bbox_y0"]], "bbox_y0"),
            bbox_x1=parse_float(excel_row[header_map["bbox_x1"]], "bbox_x1"),
            bbox_y1=parse_float(excel_row[header_map["bbox_y1"]], "bbox_y1"),
            font=str(excel_row[header_map["font"]] or ""),
            font_size=parse_float(excel_row[header_map["font_size"]], "font_size"),
            flags=parse_int(excel_row[header_map["flags"]], "flags"),
            block_number=parse_int(excel_row[header_map["block_number"]], "block_number"),
            line_number=parse_int(excel_row[header_map["line_number"]], "line_number"),
            origin_x=(
                parse_float(excel_row[header_map["origin_x"]], "origin_x")
                if "origin_x" in header_map
                else parse_float(excel_row[header_map["bbox_x0"]], "bbox_x0")
            ),
            origin_y=(
                parse_float(excel_row[header_map["origin_y"]], "origin_y")
                if "origin_y" in header_map
                else parse_float(excel_row[header_map["bbox_y1"]], "bbox_y1")
            ),
        )
        span_map[span_id] = span

    return span_map


def read_translation_data(excel_path: Path, sheet_name: str) -> tuple[list[TranslationUnitRow], dict[int, TextSpanRow]]:
    try:
        workbook = load_workbook(excel_path, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"Could not open Excel file: {excel_path} | {exc}") from exc

    rows = read_translation_units(excel_path, sheet_name, workbook=workbook)
    span_map = read_text_spans(workbook)
    return rows, span_map


def get_unit_spans(row: TranslationUnitRow, span_map: dict[int, TextSpanRow]) -> list[TextSpanRow]:
    spans: list[TextSpanRow] = []
    for span_id in parse_source_span_ids(row.source_span_ids):
        span = span_map.get(span_id)
        if span is None:
            continue
        if span.page != row.page:
            continue
        spans.append(span)

    spans.sort(key=lambda item: (item.block_number, item.line_number, item.origin_y, item.origin_x))
    return spans


def compute_lineheight_from_spans(spans: list[TextSpanRow], font_size: float) -> float | None:
    if not spans:
        return None

    origins: list[float] = []
    seen: set[tuple[int, int, float]] = set()
    for span in spans:
        key = (span.block_number, span.line_number, round(span.origin_y, 2))
        if key in seen:
            continue
        seen.add(key)
        origins.append(span.origin_y)

    origins.sort()
    if len(origins) < 2:
        return None

    deltas = [origins[index + 1] - origins[index] for index in range(len(origins) - 1) if origins[index + 1] > origins[index]]
    if not deltas:
        return None

    median_delta = statistics.median(deltas)
    base = max(0.1, font_size)
    ratio = max(1.0, min(1.8, median_delta / base))
    return round(ratio, 3)


def is_table_header_candidate(unit: TranslationUnitRow) -> bool:
    return is_page2_table_header_row(unit) or is_page3_table_header_row(unit)


def is_explicit_center_row(unit: TranslationUnitRow) -> bool:
    normalized = normalize_text(display_text_for_row(unit))
    if normalized in EXPLICIT_CENTER_TEXTS:
        return True
    if is_page2_table_header_row(unit) or is_page3_table_header_row(unit):
        return True
    return False


def is_centered_unit(unit: TranslationUnitRow, page_rect: fitz.Rect) -> bool:
    _ = page_rect
    return is_explicit_center_row(unit)


def should_expand_text_box(unit: TranslationUnitRow) -> bool:
    text = unit.original_text.strip()
    return is_table_header_candidate(unit) and len(text) <= 30


def build_text_rect(unit: TranslationUnitRow) -> fitz.Rect:
    rect = fitz.Rect(unit.bbox)
    if should_expand_text_box(unit):
        rect = fitz.Rect(rect.x0 - 2.0, rect.y0, rect.x1 + 2.0, rect.y1)
    return rect


def min_font_size_for(unit: TranslationUnitRow) -> float:
    if unit.font_size < 5.0:
        return 3.2
    return 4.0


def alignment_for(unit: TranslationUnitRow, page_rect: fitz.Rect) -> int:
    _ = page_rect
    if is_centered_unit(unit, page_rect):
        return fitz.TEXT_ALIGN_CENTER
    return fitz.TEXT_ALIGN_LEFT


def forced_table_text(row: TranslationUnitRow) -> str | None:
    key = normalize_text(display_text_for_row(row))
    if row.page == 2:
        return PAGE2_TABLE_HEADER_TEXTS.get(key)
    if row.page == 3:
        return PAGE3_TABLE_HEADER_TEXTS.get(key)
    return None


def estimate_text_width(text: str, font_spec: FontSpec, font_size: float) -> float:
    longest = max(text.split("\n"), key=len, default="")
    try:
        return fitz.get_text_length(longest, fontname=font_spec.fontname, fontsize=font_size)
    except Exception:  # noqa: BLE001
        return len(longest) * font_size * 0.55


def evaluate_text_fit(
    page_rect: fitz.Rect,
    unit: TranslationUnitRow,
    font_spec: FontSpec,
    *,
    text_override: str | None = None,
    font_size: float | None = None,
    align_override: int | None = None,
    lineheight: float | None = None,
) -> FitResult:
    temp_doc = fitz.open()
    try:
        temp_page = temp_doc.new_page(width=page_rect.width, height=page_rect.height)
        text_rect = build_text_rect(unit)
        text_value = text_override if text_override is not None else unit.english_translation
        original_size = max(0.1, unit.font_size)
        current_size = original_size if font_size is None else max(0.1, font_size)
        min_size = min_font_size_for(unit)
        reduction_step = 0.25
        align = alignment_for(unit, page_rect) if align_override is None else align_override
        allow_reduce = font_size is None
        has_more_than_100 = len(text_value) > 100

        while current_size >= min_size:
            result = temp_page.insert_textbox(
                text_rect,
                text_value,
                fontsize=current_size,
                fontname=font_spec.fontname,
                fontfile=font_spec.fontfile,
                color=fitz.sRGB_to_pdf(unit.color),
                align=align,
                lineheight=lineheight,
            )
            if result >= 0:
                return FitResult(fits=True, proposed_size=round(current_size, 2))

            if not allow_reduce:
                break
            if has_more_than_100 and current_size > original_size * 0.9:
                current_size = max(original_size * 0.9, current_size - reduction_step)
            else:
                current_size -= reduction_step

        return FitResult(fits=False, proposed_size=round(max(current_size, min_size), 2))
    finally:
        temp_doc.close()


def evaluate_fit_in_rect(
    rect: fitz.Rect,
    text: str,
    unit: TranslationUnitRow,
    font_spec: FontSpec,
    font_size: float,
    align: int,
    lineheight: float | None = None,
) -> bool:
    temp_doc = fitz.open()
    try:
        temp_page = temp_doc.new_page(width=max(rect.x1 + 10.0, 100.0), height=max(rect.y1 + 10.0, 100.0))
        result = temp_page.insert_textbox(
            rect,
            text,
            fontsize=font_size,
            fontname=font_spec.fontname,
            fontfile=font_spec.fontfile,
            color=fitz.sRGB_to_pdf(unit.color),
            align=align,
            lineheight=lineheight,
        )
        return result >= 0
    finally:
        temp_doc.close()


def collect_arial_paths() -> dict[str, Path]:
    return {
        "arial": Path(r"C:\Windows\Fonts\arial.ttf"),
        "arialbd": Path(r"C:\Windows\Fonts\arialbd.ttf"),
        "arialbi": Path(r"C:\Windows\Fonts\arialbi.ttf"),
        "ariali": Path(r"C:\Windows\Fonts\ariali.ttf"),
    }


def build_default_spec(row: TranslationUnitRow) -> PlacementSpec:
    mode = "TEXTBOX"
    use_textbox = True
    align = fitz.TEXT_ALIGN_LEFT

    if is_explicit_center_row(row):
        align = fitz.TEXT_ALIGN_CENTER

    if is_short_label_row(row) or is_section_title_row(row):
        mode = "BASELINE"
        use_textbox = False
        if is_explicit_center_row(row):
            align = fitz.TEXT_ALIGN_CENTER
        else:
            align = fitz.TEXT_ALIGN_LEFT

    if is_page2_table_header_row(row) or is_page3_table_header_row(row):
        mode = "TABLE_GROUP"
        use_textbox = True
        align = fitz.TEXT_ALIGN_CENTER

    if is_notes_row(row):
        mode = "NOTES_GROUP"
        use_textbox = True
        align = fitz.TEXT_ALIGN_LEFT

    return PlacementSpec(
        placement_mode=mode,
        font_size_original=round(max(0.1, row.font_size), 2),
        font_size_final=round(max(0.1, row.font_size), 2),
        align=align,
        use_textbox=use_textbox,
    )


def plan_page2_paragraph_group(
    page_rect: fitz.Rect,
    paragraph_rows: list[TranslationUnitRow],
    specs: dict[int, PlacementSpec],
    span_map: dict[int, TextSpanRow],
    arial_paths: dict[str, Path],
) -> None:
    if not paragraph_rows:
        return

    size_candidates = [max(0.1, row.font_size) for row in paragraph_rows]
    common_size = round(statistics.median(size_candidates), 2)

    lineheight_values: list[float] = []
    for row in paragraph_rows:
        lineheight = compute_lineheight_from_spans(get_unit_spans(row, span_map), row.font_size)
        if lineheight is not None:
            lineheight_values.append(lineheight)
    common_lineheight = round(statistics.median(lineheight_values), 3) if lineheight_values else 1.2

    while common_size >= 4.0:
        all_fit = True
        for row in paragraph_rows:
            font_spec = choose_font_spec(row, arial_paths)
            fit_result = evaluate_text_fit(
                page_rect,
                row,
                font_spec,
                font_size=common_size,
                align_override=fitz.TEXT_ALIGN_LEFT,
                lineheight=common_lineheight,
            )
            if not fit_result.fits:
                all_fit = False
                break
        if all_fit:
            break
        common_size = round(common_size - 0.25, 2)

    for row in paragraph_rows:
        spec = specs[row.unit_id]
        spec.font_size_final = common_size
        spec.align = fitz.TEXT_ALIGN_LEFT
        spec.use_textbox = True
        spec.lineheight = common_lineheight
        spec.placement_mode = "TEXTBOX"


def plan_table_group(
    page_rect: fitz.Rect,
    table_rows: list[TranslationUnitRow],
    specs: dict[int, PlacementSpec],
    arial_paths: dict[str, Path],
) -> None:
    if not table_rows:
        return

    size_candidates = [max(0.1, row.font_size) for row in table_rows]
    common_size = round(statistics.median(size_candidates), 2)

    while common_size >= 4.0:
        all_fit = True
        for row in table_rows:
            forced = forced_table_text(row) or row.english_translation
            font_spec = choose_font_spec(row, arial_paths)
            fit = evaluate_text_fit(
                page_rect,
                row,
                font_spec,
                text_override=forced,
                font_size=common_size,
                align_override=fitz.TEXT_ALIGN_CENTER,
            )
            if not fit.fits:
                all_fit = False
                break
        if all_fit:
            break
        common_size = round(common_size - 0.25, 2)

    for row in table_rows:
        spec = specs[row.unit_id]
        spec.font_size_final = common_size
        spec.align = fitz.TEXT_ALIGN_CENTER
        spec.use_textbox = True
        spec.placement_mode = "TABLE_GROUP"
        spec.forced_text = forced_table_text(row)


def compute_notes_zone(notes_rows: list[TranslationUnitRow], observations_row: TranslationUnitRow) -> fitz.Rect:
    notes_sorted = sorted(notes_rows, key=lambda row: row.bbox_y0)
    left = notes_sorted[0].bbox_x0
    right = max(row.bbox_x1 for row in notes_sorted)
    top = notes_sorted[0].bbox_y0
    bottom = observations_row.bbox_y0
    return fitz.Rect(left, top, right, bottom)


def compose_notes_text(notes_rows: list[TranslationUnitRow]) -> str:
    notes_sorted = sorted(notes_rows, key=lambda row: row.bbox_y0)
    return "\n\n".join(row.english_translation.strip() for row in notes_sorted)


def plan_notes_group(
    notes_rows: list[TranslationUnitRow],
    observations_row: TranslationUnitRow | None,
    specs: dict[int, PlacementSpec],
    arial_paths: dict[str, Path],
) -> tuple[float | None, bool]:
    if not notes_rows or observations_row is None:
        return None, False

    common_size = round(statistics.median([max(0.1, row.font_size) for row in notes_rows]), 2)
    min_size = 4.0
    zone = compute_notes_zone(notes_rows, observations_row)
    text = compose_notes_text(notes_rows)
    sample_row = notes_rows[0]
    font_spec = choose_font_spec(sample_row, arial_paths)

    while common_size >= min_size:
        if evaluate_fit_in_rect(zone, text, sample_row, font_spec, common_size, fitz.TEXT_ALIGN_LEFT, lineheight=1.2):
            break
        common_size = round(common_size - 0.25, 2)

    fits = common_size >= min_size and evaluate_fit_in_rect(
        zone,
        text,
        sample_row,
        font_spec,
        common_size,
        fitz.TEXT_ALIGN_LEFT,
        lineheight=1.2,
    )

    for row in notes_rows:
        spec = specs[row.unit_id]
        spec.font_size_final = common_size
        spec.align = fitz.TEXT_ALIGN_LEFT
        spec.use_textbox = True
        spec.placement_mode = "NOTES_GROUP"
        spec.lineheight = 1.2

    return common_size, fits


def build_page_specs(
    page_rect: fitz.Rect,
    page_rows: list[TranslationUnitRow],
    span_map: dict[int, TextSpanRow],
    arial_paths: dict[str, Path],
) -> tuple[dict[int, PlacementSpec], dict[str, Any]]:
    specs: dict[int, PlacementSpec] = {}
    warnings: list[str] = []
    failures: list[str] = []

    translate_rows = [row for row in page_rows if row.action == "TRANSLATE" and row.english_translation.strip() != ""]
    for row in translate_rows:
        specs[row.unit_id] = build_default_spec(row)

    page2_paragraph_rows = [row for row in translate_rows if is_page2_paragraph_row(row)]
    plan_page2_paragraph_group(page_rect, page2_paragraph_rows, specs, span_map, arial_paths)

    page2_table_rows = [row for row in translate_rows if is_page2_table_header_row(row)]
    page3_table_rows = [row for row in translate_rows if is_page3_table_header_row(row)]
    plan_table_group(page_rect, page2_table_rows, specs, arial_paths)
    plan_table_group(page_rect, page3_table_rows, specs, arial_paths)

    notes_rows = [row for row in translate_rows if is_notes_row(row)]
    observations_row = next((row for row in translate_rows if normalize_text(display_text_for_row(row)) == "observations"), None)
    notes_size, notes_fit = plan_notes_group(notes_rows, observations_row, specs, arial_paths)
    if notes_rows and not notes_fit:
        failures.append("NOTES_GROUP_OVERFLOW")

    for row in page2_table_rows + page3_table_rows:
        spec = specs[row.unit_id]
        font_spec = choose_font_spec(row, arial_paths)
        text = spec.forced_text if spec.forced_text is not None else row.english_translation
        width = estimate_text_width(text, font_spec, spec.font_size_final)
        if width > (row.bbox_x1 - row.bbox_x0) + 0.1:
            failures.append(f"TABLE_CELL_CROSS:{row.unit_id}")
        if spec.font_size_final < max(0.1, row.font_size) * 0.85:
            warnings.append(f"TABLE_SIZE_BELOW_85_PERCENT:{row.unit_id}")

    if notes_rows and observations_row is not None:
        zone = compute_notes_zone(notes_rows, observations_row)
        if zone.y1 <= zone.y0:
            failures.append("NOTES_GROUP_INVALID_ZONE")
        if notes_size is not None and notes_size < max(0.1, statistics.median([row.font_size for row in notes_rows])) * 0.85:
            warnings.append("NOTES_SIZE_BELOW_85_PERCENT")

    return specs, {"warnings": warnings, "failures": failures}


def preflight_report_path(project_root: Path, output_path: Path | None, source_pdf: Path) -> Path:
    output_dir = project_root / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        stem = f"{source_pdf.stem}_EN_TRANSLATION"
    else:
        resolved_output = output_path if output_path.is_absolute() else (project_root / output_path)
        stem = resolved_output.stem
    return (output_dir / f"{stem}_preflight_report.json").resolve()


def run_preflight(
    doc: fitz.Document,
    rows: list[TranslationUnitRow],
    span_map: dict[int, TextSpanRow],
) -> dict[str, Any]:
    arial_paths = collect_arial_paths()
    rows_by_page = group_rows_by_page(rows)

    entries: list[dict[str, Any]] = []
    failed_pages: set[int] = set()

    for page_index in range(doc.page_count):
        page_number = page_index + 1
        page = doc.load_page(page_index)
        page_rows = rows_by_page.get(page_number, [])
        specs, page_meta = build_page_specs(page.rect, page_rows, span_map, arial_paths)

        if page_meta["failures"]:
            failed_pages.add(page_number)

        for row in page_rows:
            if row.action != "TRANSLATE":
                continue

            spec = specs.get(row.unit_id, build_default_spec(row))
            text_value = spec.forced_text if spec.forced_text is not None else row.english_translation
            fits = False

            if row.page < 1 or row.page > doc.page_count:
                fits = False
            elif row.english_translation.strip() == "":
                fits = False
            elif spec.placement_mode == "BASELINE":
                spans = get_unit_spans(row, span_map)
                if "\n" not in text_value and len(text_value) < 80:
                    first_span = spans[0] if spans else None
                    if first_span is not None:
                        font_spec = choose_font_spec(row, arial_paths)
                        width = estimate_text_width(text_value, font_spec, spec.font_size_final)
                        source_width = max(1.0, first_span.bbox_x1 - first_span.bbox_x0)
                        fits = width <= source_width + 0.1
                    else:
                        fits = False
                else:
                    font_spec = choose_font_spec(row, arial_paths)
                    fit = evaluate_text_fit(
                        page.rect,
                        row,
                        font_spec,
                        text_override=text_value,
                        font_size=spec.font_size_final,
                        align_override=spec.align,
                    )
                    fits = fit.fits
            elif spec.placement_mode == "NOTES_GROUP":
                fits = "NOTES_GROUP_OVERFLOW" not in page_meta["failures"]
            else:
                font_spec = choose_font_spec(row, arial_paths)
                fit = evaluate_text_fit(
                    page.rect,
                    row,
                    font_spec,
                    text_override=text_value,
                    font_size=spec.font_size_final,
                    align_override=spec.align,
                    lineheight=spec.lineheight,
                )
                fits = fit.fits

            if not fits:
                failed_pages.add(page_number)

            entries.append(
                {
                    "unit_id": row.unit_id,
                    "page": row.page,
                    "text": row.english_translation,
                    "proposed_size": spec.font_size_final,
                    "font_size_original": spec.font_size_original,
                    "font_size_final": spec.font_size_final,
                    "placement_mode": spec.placement_mode,
                    "fits": fits,
                }
            )

    page_status: dict[str, str] = {}
    for page_index in range(1, doc.page_count + 1):
        page_status[str(page_index)] = "FAILED_PREFLIGHT" if page_index in failed_pages else "OK"

    return {
        "status": "FAILED_PREFLIGHT" if failed_pages else "OK",
        "failed_pages": sorted(failed_pages),
        "page_status": page_status,
        "units": entries,
    }


def save_preflight_report(report_path: Path, report: dict[str, Any]) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)


def try_insert_translation(
    page: fitz.Page,
    unit: TranslationUnitRow,
    font_spec: FontSpec,
    summary: Summary,
    spec: PlacementSpec,
    spans: list[TextSpanRow],
) -> bool:
    text_value = spec.forced_text if spec.forced_text is not None else unit.english_translation
    fontsize = max(0.1, spec.font_size_final)

    if spec.placement_mode == "BASELINE":
        if not spans:
            summary.not_fit_texts += 1
            LOGGER.warning("Missing source spans for baseline placement | unit_id=%s", unit.unit_id)
            return False

        color = fitz.sRGB_to_pdf(unit.color)
        lines = text_value.splitlines() if "\n" in text_value else [text_value]
        first_span = spans[0]
        origin_x = first_span.origin_x

        if len(lines) == 1:
            text_width = estimate_text_width(lines[0], font_spec, fontsize)
            span_width = max(1.0, first_span.bbox_x1 - first_span.bbox_x0)
            if text_width > span_width + 0.1:
                summary.not_fit_texts += 1
                return False
            if spec.align == fitz.TEXT_ALIGN_CENTER:
                origin_x = unit.bbox_x0 + ((unit.bbox_x1 - unit.bbox_x0) - text_width) / 2.0
            page.insert_text(
                point=(origin_x, first_span.origin_y),
                text=lines[0],
                fontsize=fontsize,
                fontname=font_spec.fontname,
                fontfile=font_spec.fontfile,
                color=color,
                render_mode=0,
            )
            return True

        line_origins: list[tuple[float, float]] = []
        seen_lines: set[tuple[int, int]] = set()
        for span in spans:
            key = (span.block_number, span.line_number)
            if key in seen_lines:
                continue
            seen_lines.add(key)
            line_origins.append((span.origin_x, span.origin_y))
        line_origins.sort(key=lambda item: item[1])

        spacing = fontsize * 1.2
        if len(line_origins) > 1:
            deltas = [line_origins[index + 1][1] - line_origins[index][1] for index in range(len(line_origins) - 1)]
            positive = [delta for delta in deltas if delta > 0]
            if positive:
                spacing = statistics.median(positive)

        for index, text_line in enumerate(lines):
            if index < len(line_origins):
                line_x, line_y = line_origins[index]
            else:
                line_x = line_origins[-1][0]
                line_y = line_origins[-1][1] + spacing * (index - len(line_origins) + 1)
            page.insert_text(
                point=(line_x, line_y),
                text=text_line,
                fontsize=fontsize,
                fontname=font_spec.fontname,
                fontfile=font_spec.fontfile,
                color=color,
                render_mode=0,
            )
        return True

    text_rect = build_text_rect(unit)
    result = page.insert_textbox(
        text_rect,
        text_value,
        fontsize=fontsize,
        fontname=font_spec.fontname,
        fontfile=font_spec.fontfile,
        color=fitz.sRGB_to_pdf(unit.color),
        align=spec.align,
        lineheight=spec.lineheight,
    )
    if result >= 0:
        return True

    summary.not_fit_texts += 1
    LOGGER.warning("Text does not fit | unit_id=%s | page=%s | text=%s", unit.unit_id, unit.page, text_value)
    return False


def add_redaction(page: fitz.Page, unit: TranslationUnitRow) -> None:
    rect = fitz.Rect(unit.bbox)
    redaction_rect = fitz.Rect(rect.x0 - 0.4, rect.y0 - 0.4, rect.x1 + 0.4, rect.y1 + 0.4)
    page.add_redact_annot(redaction_rect, fill=None, cross_out=False)


def note_rect_candidate(page: fitz.Page) -> fitz.Rect:
    rect = page.rect
    margin_x = 18.0
    note_height = 6.0
    y1 = rect.y1 - 2.0
    y0 = y1 - note_height
    return fitz.Rect(rect.x0 + margin_x, y0, rect.x1 - margin_x, y1)


def has_space_for_note(page: fitz.Page, rect: fitz.Rect) -> bool:
    text_dict = page.get_text("dict")
    for block in text_dict.get("blocks", []):
        block_bbox = block.get("bbox")
        if not isinstance(block_bbox, (list, tuple)) or len(block_bbox) != 4:
            continue
        existing = fitz.Rect(float(block_bbox[0]), float(block_bbox[1]), float(block_bbox[2]), float(block_bbox[3]))
        if existing.intersects(rect):
            return False
    return True


def add_translation_note(page: fitz.Page) -> bool:
    rect = note_rect_candidate(page)
    if not has_space_for_note(page, rect):
        LOGGER.warning("Skipping translation note due to insufficient free space | page=%s", page.number + 1)
        return False

    result = page.insert_textbox(
        rect,
        TRANSLATION_NOTE,
        fontsize=4.0,
        fontname="helv",
        color=fitz.sRGB_to_pdf(0x555555),
        align=fitz.TEXT_ALIGN_CENTER,
    )
    if result < 0:
        LOGGER.warning("Failed to place translation note | page=%s", page.number + 1)
        return False
    return True


def ensure_output_paths(project_root: Path, output_path: Path | None, source_pdf: Path) -> tuple[Path, Path]:
    output_dir = project_root / "output"
    renders_dir = project_root / "renders"
    output_dir.mkdir(parents=True, exist_ok=True)
    renders_dir.mkdir(parents=True, exist_ok=True)

    if output_path is None:
        final_pdf_path = output_dir / f"{source_pdf.stem}_EN_TRANSLATION.pdf"
    else:
        final_pdf_path = output_path if output_path.is_absolute() else (project_root / output_path)

    final_pdf_path.parent.mkdir(parents=True, exist_ok=True)
    render_root = renders_dir / final_pdf_path.stem
    render_root.mkdir(parents=True, exist_ok=True)
    return final_pdf_path.resolve(), render_root.resolve()


def group_rows_by_page(rows: list[TranslationUnitRow]) -> dict[int, list[TranslationUnitRow]]:
    grouped: dict[int, list[TranslationUnitRow]] = {}
    for row in rows:
        grouped.setdefault(row.page, []).append(row)
    for page_rows in grouped.values():
        page_rows.sort(key=lambda row: (row.bbox_y0, row.bbox_x0, row.unit_id))
    return grouped


def process_pdf(
    doc: fitz.Document,
    rows: list[TranslationUnitRow],
    span_map: dict[int, TextSpanRow],
    summary: Summary,
) -> None:
    arial_paths = collect_arial_paths()

    rows_by_page = group_rows_by_page(rows)

    for page_index in range(doc.page_count):
        page_number = page_index + 1
        page = doc.load_page(page_index)
        page_rows = rows_by_page.get(page_number, [])
        translatable_rows: list[TranslationUnitRow] = []
        specs, _ = build_page_specs(page.rect, page_rows, span_map, arial_paths)
        notes_rows = [row for row in page_rows if row.action == "TRANSLATE" and is_notes_row(row)]
        observations_row = next(
            (
                row
                for row in page_rows
                if row.action == "TRANSLATE" and normalize_text(display_text_for_row(row)) == "observations"
            ),
            None,
        )

        for row in page_rows:
            try:
                if row.action == "KEEP":
                    summary.keep_ignored += 1
                    continue
                if row.action not in {"TRANSLATE", "REMOVE"}:
                    continue

                if row.page < 1 or row.page > doc.page_count:
                    raise ValueError("Page index out of range")

                add_redaction(page, row)
                if row.action == "REMOVE":
                    summary.remove_processed += 1
                    continue

                if row.english_translation.strip() == "":
                    summary.empty_translations += 1
                    LOGGER.warning(
                        "Empty translation; skipping insertion | unit_id=%s | page=%s | original=%s",
                        row.unit_id,
                        row.page,
                        row.original_text,
                    )
                    continue

                if is_notes_row(row):
                    # Notes are composed and inserted as a single group later on page 3.
                    translatable_rows.append(row)
                    summary.translate_processed += 1
                    continue

                translatable_rows.append(row)
                summary.translate_processed += 1

            except Exception as exc:  # noqa: BLE001
                summary.row_failures += 1
                LOGGER.error(
                    "Row processing failed before redaction apply | unit_id=%s | page=%s | text=%s | error=%s",
                    row.unit_id,
                    row.page,
                    row.original_text,
                    exc,
                )

        if page.first_annot is not None:
            page.apply_redactions(
                images=fitz.PDF_REDACT_IMAGE_NONE,
                graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                text=fitz.PDF_REDACT_TEXT_REMOVE,
            )

        for row in translatable_rows:
            if is_notes_row(row):
                continue
            try:
                font_spec = choose_font_spec(row, arial_paths)
                spec = specs.get(row.unit_id, build_default_spec(row))
                spans = get_unit_spans(row, span_map)
                inserted = try_insert_translation(page, row, font_spec, summary, spec, spans)
                if not inserted:
                    raise RuntimeError(
                        f"Unexpected insertion failure after redaction | unit_id={row.unit_id} | page={row.page}"
                    )
            except Exception as exc:  # noqa: BLE001
                summary.row_failures += 1
                raise RuntimeError(
                    f"Row insertion failed | unit_id={row.unit_id} | page={row.page} | text={row.original_text} | error={exc}"
                ) from exc

        if notes_rows and observations_row is not None:
            try:
                zone = compute_notes_zone(notes_rows, observations_row)
                notes_text = compose_notes_text(notes_rows)
                first_note = sorted(notes_rows, key=lambda item: item.bbox_y0)[0]
                spec = specs.get(first_note.unit_id, build_default_spec(first_note))
                font_spec = choose_font_spec(first_note, arial_paths)
                result = page.insert_textbox(
                    zone,
                    notes_text,
                    fontsize=spec.font_size_final,
                    fontname=font_spec.fontname,
                    fontfile=font_spec.fontfile,
                    color=fitz.sRGB_to_pdf(first_note.color),
                    align=fitz.TEXT_ALIGN_LEFT,
                    lineheight=spec.lineheight,
                )
                if result < 0:
                    raise RuntimeError("Notes group insertion failed")
            except Exception as exc:  # noqa: BLE001
                summary.row_failures += 1
                raise RuntimeError(f"Notes group insertion failed | page={page_number} | error={exc}") from exc

        add_translation_note(page)


def render_validation_images(doc_path: Path, render_root: Path) -> None:
    dpi = 150
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    with fitz.open(doc_path) as doc:
        for index in range(doc.page_count):
            page = doc.load_page(index)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            pix.save(render_root / f"page_{index + 1:03d}.png")


def print_summary(summary: Summary, output_pdf: Path, render_root: Path) -> None:
    print("Translation apply summary")
    print(f"- TRANSLATE units processed: {summary.translate_processed}")
    print(f"- REMOVE units processed: {summary.remove_processed}")
    print(f"- KEEP units ignored: {summary.keep_ignored}")
    print(f"- Empty translations: {summary.empty_translations}")
    print(f"- Texts with reduced font: {summary.reduced_font_texts}")
    print(f"- Texts that did not fit: {summary.not_fit_texts}")
    print(f"- Row failures: {summary.row_failures}")
    print(f"- Output PDF: {output_pdf}")
    print(f"- Validation renders: {render_root}")


def main() -> int:
    setup_logging()
    args = parse_args()

    script_path = Path(__file__).resolve()
    project_root = script_path.parent.parent
    pdf_path = to_project_path(project_root, args.pdf_path)
    excel_path = to_project_path(project_root, args.excel_path)

    if not pdf_path.exists():
        LOGGER.error("PDF not found: %s", pdf_path)
        return 1
    if not excel_path.exists():
        LOGGER.error("Excel not found: %s", excel_path)
        return 1

    report_path = preflight_report_path(project_root, args.output, pdf_path)

    try:
        rows, span_map = read_translation_data(excel_path, args.sheet)
    except Exception as exc:  # noqa: BLE001
        LOGGER.error("Failed to read translation sheet: %s", exc)
        return 1

    try:
        with fitz.open(pdf_path) as preflight_doc:
            report = run_preflight(preflight_doc, rows, span_map)
    except Exception as exc:  # noqa: BLE001
        LOGGER.error("Failed during preflight: %s", exc)
        return 1

    save_preflight_report(report_path, report)
    if args.dry_run:
        LOGGER.info("Dry-run completed. Preflight report: %s", report_path)
        if report["status"] != "OK":
            return 2
        return 0

    if report["status"] != "OK":
        LOGGER.error("Preflight failed on pages: %s", report["failed_pages"])
        LOGGER.error("PDF generation cancelled. Report: %s", report_path)
        return 2

    output_pdf, render_root = ensure_output_paths(project_root, args.output, pdf_path)

    try:
        with fitz.open(pdf_path) as doc:
            summary = Summary()
            process_pdf(doc, rows, span_map, summary)
            doc.save(output_pdf, garbage=4, deflate=True, clean=True)
    except Exception as exc:  # noqa: BLE001
        LOGGER.error("Failed to process PDF: %s", exc)
        return 1

    try:
        render_validation_images(output_pdf, render_root)
    except Exception as exc:  # noqa: BLE001
        LOGGER.error("Failed to render validation images: %s", exc)
        return 1

    print_summary(summary, output_pdf, render_root)
    return 0


if __name__ == "__main__":
    sys.exit(main())
